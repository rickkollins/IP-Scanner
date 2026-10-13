"""Write a scanned book to a Word document (.docx).

Every book page becomes its own section that starts on a new page, and the
section's page numbering restarts at the number printed in the book. The
footer holds a PAGE field, so Word shows the book's own page numbers - and
skipped or unnumbered pages keep the right numbers too.
"""

from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from book import Book

# Elements that must come after w:pgNumType inside w:sectPr (ECMA-376 order).
_AFTER_PGNUMTYPE = ("cols", "formProt", "vAlign", "noEndnote", "titlePg",
                    "textDirection", "bidi", "rtlGutter", "docGrid",
                    "printerSettings", "sectPrChange")


def _set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        for tag in _AFTER_PGNUMTYPE:
            follower = sect_pr.find(qn(f"w:{tag}"))
            if follower is not None:
                follower.addprevious(pg)
                break
        else:
            sect_pr.append(pg)
    pg.set(qn("w:start"), str(start))


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None),
                       (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == " PAGE ":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


PAPER_SIZES = {  # name: (width, height, side margin) in inches
    "Letter (8.5 × 11 in)": (8.5, 11.0, 1.0),
    "A4 (210 × 297 mm)": (8.27, 11.69, 1.0),
    "Book (6 × 9 in)": (6.0, 9.0, 0.75),
}


def export(book: Book, path: str, paper: str = "Letter (8.5 × 11 in)") -> int:
    """Save `book` as a .docx at `path`. Returns the number of pages written."""
    width, height, margin = PAPER_SIZES.get(paper, next(iter(PAPER_SIZES.values())))
    book.renumber()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    doc.core_properties.title = book.title

    written = 0
    for i, page in enumerate(book.pages):
        if i == 0:
            section = doc.sections[0]
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.page_width = Inches(width)
        section.page_height = Inches(height)
        section.left_margin = section.right_margin = Inches(margin)
        section.top_margin = section.bottom_margin = Inches(margin)
        _set_page_number_start(section, page.number)

        paragraphs = page.paragraphs()
        if not paragraphs:
            doc.add_paragraph()  # keep blank pages so the numbering lines up
        for para in paragraphs:
            if para.heading:
                doc.add_heading(para.text, level=2)
            else:
                doc.add_paragraph(para.text)
        written += 1

    # All sections share the first section's footer; its PAGE field shows
    # each section's own (restarted) number.
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_field(fp)
    doc.save(path)
    return written
